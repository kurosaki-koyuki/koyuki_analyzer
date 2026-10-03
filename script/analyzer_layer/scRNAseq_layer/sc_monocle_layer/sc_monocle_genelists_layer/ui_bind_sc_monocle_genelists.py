# -*- coding: utf-8 -*-
"""
scRNAseq Monocle基因列表分析子层功能绑定脚本
负责：1) 绑定音乐控制器；2) 绑定运行按钮、阈值筛选、基因搜索等；3) QThread后台执行graph_test
"""

import os
import traceback

from PyQt5.QtCore import QThread, pyqtSignal, QObject

from script.utils_layer.music_controller_fix import fix_music_controller_bindings
from script.analyzer_layer.scRNAseq_layer.sc_monocle_layer.sc_monocle_genelists_layer.sc_monocle_genelists_analysis import ScMonocleGenelistsAnalysis
from script.analyzer_layer.scRNAseq_layer.sc_monocle_layer.sc_monocle_genelists_layer.ui_func_sc_monocle_genelists import ScMonocleGenelistsFunc


class GraphTestWorker(QObject):
    """graph_test后台执行Worker（在QThread中运行subprocess）"""
    finished = pyqtSignal(bool, object)  # (success, result_df_or_error_str)
    progress = pyqtSignal(str)

    def __init__(self, analysis, gene_filter_mode, n_top, gene_list_file):
        super().__init__()
        self.analysis = analysis
        self.gene_filter_mode = gene_filter_mode
        self.n_top = n_top
        self.gene_list_file = gene_list_file
        self._running = True

    def stop(self):
        self._running = False

    def run(self):
        try:
            self.progress.emit("开始运行graph_test...")

            success, result, is_error = self.analysis.run_graph_test(
                gene_filter_mode=self.gene_filter_mode,
                n_top=self.n_top,
                gene_list_file=self.gene_list_file,
                progress_callback=self._on_progress
            )

            if success:
                self.progress.emit(f"graph_test完成, 结果行数: {len(result)}")
                self.finished.emit(True, result)
            else:
                self.progress.emit(f"graph_test失败: {result}")
                self.finished.emit(False, result)
        except Exception as e:
            err_msg = f"graph_test执行异常: {e}\n{traceback.format_exc()}"
            self.progress.emit(err_msg)
            self.finished.emit(False, str(e))
        finally:
            self._running = False

    def _on_progress(self, msg):
        """subprocess进度回调"""
        if not self._running:
            return
        self.progress.emit(msg)


class SpearmanWorker(QObject):
    """阶段三基因表达与伪时间关系分析后台执行Worker（在QThread中运行subprocess）"""
    finished = pyqtSignal(bool, object)  # (success, result_df_or_error_str)
    progress = pyqtSignal(str)

    def __init__(self, analysis, calc_scope, threshold_field, threshold_value, algorithm):
        super().__init__()
        self.analysis = analysis
        self.calc_scope = calc_scope
        self.threshold_field = threshold_field
        self.threshold_value = threshold_value
        self.algorithm = algorithm
        self._running = True

    def stop(self):
        self._running = False

    def run(self):
        try:
            self.progress.emit(f"开始运行{self.algorithm}分析...")

            success, result, is_error = self.analysis.run_spearman_analysis(
                calc_scope=self.calc_scope,
                threshold_field=self.threshold_field,
                threshold_value=self.threshold_value,
                algorithm=self.algorithm,
                progress_callback=self._on_progress
            )

            if success:
                self.progress.emit(f"{self.algorithm}分析完成, 结果行数: {len(result)}")
                self.finished.emit(True, result)
            else:
                self.progress.emit(f"{self.algorithm}分析失败: {result}")
                self.finished.emit(False, result)
        except Exception as e:
            err_msg = f"{self.algorithm}分析执行异常: {e}\n{traceback.format_exc()}"
            self.progress.emit(err_msg)
            self.finished.emit(False, str(e))
        finally:
            self._running = False

    def _on_progress(self, msg):
        """subprocess进度回调"""
        if not self._running:
            return
        self.progress.emit(msg)


class ScMonocleGenelistsBind:
    """Monocle基因列表类功能绑定"""

    def __init__(self, main_window, genelists_ui):
        self.parent = main_window
        self.sc_monocle_genelists_ui = genelists_ui
        self.analysis = ScMonocleGenelistsAnalysis()
        self.func = ScMonocleGenelistsFunc(genelists_ui, main_window)
        self._worker = None
        self._thread = None
        self._spearman_worker = None  # 阶段三Worker
        self._spearman_thread = None  # 阶段三Thread
        self._imported_gene_list_file = None  # 缓存外部基因列表文件路径

        # 初始化基因搜索（TableSearcherMixin）
        # 表格列表顺序必须与标签页顺序一致：
        # 0:所有基因列表 1:显著基因列表 2:上下调总体 3:上调基因 4:下调基因
        self.func.setup_gene_search(
            tab_widget=genelists_ui.genelists_tabs,
            tables=[
                genelists_ui.all_genes_table,
                genelists_ui.sig_genes_table,
                genelists_ui.stage3_all_table,
                genelists_ui.stage3_up_table,
                genelists_ui.stage3_down_table,
            ],
            gene_col=0
        )

        self.bind_signals()

        # 绑定音乐控制器（按项目规则，确保全局同步函数能发现）
        if hasattr(self.sc_monocle_genelists_ui, 'music_controller') and self.sc_monocle_genelists_ui.music_controller:
            fix_music_controller_bindings(self, self.sc_monocle_genelists_ui.music_controller)

    def bind_signals(self):
        ui = self.sc_monocle_genelists_ui

        # 基因筛选模式下拉框
        if hasattr(ui, 'gene_filter_combo'):
            ui.gene_filter_combo.currentIndexChanged.connect(self.on_filter_mode_changed)

        # 导入基因列表按钮
        if hasattr(ui, 'btn_import_gene_list'):
            ui.btn_import_gene_list.clicked.connect(self.import_gene_list)

        # 运行graph_test按钮
        if hasattr(ui, 'btn_run_graph_test'):
            ui.btn_run_graph_test.clicked.connect(self.run_graph_test)

        # 重新筛显著基因按钮
        if hasattr(ui, 'btn_apply_threshold'):
            ui.btn_apply_threshold.clicked.connect(self.apply_threshold)

        # 阶段三运行Spearman分析按钮
        if hasattr(ui, 'btn_run_stage3'):
            ui.btn_run_stage3.clicked.connect(self.run_spearman_analysis)

        # 阶段四运行火山图按钮
        if hasattr(ui, 'btn_run_stage4'):
            ui.btn_run_stage4.clicked.connect(self.run_volcano_plot)

        # 导出Excel按钮
        if hasattr(ui, 'btn_export_xlsx'):
            ui.btn_export_xlsx.clicked.connect(self.export_xlsx)

        # X轴指标切换时自动调整FC阈值
        if hasattr(ui, 'stage4_x_axis_combo') and hasattr(ui, 'stage4_fc_threshold_edit'):
            ui.stage4_x_axis_combo.currentIndexChanged.connect(self._on_x_axis_changed)

        # 基因搜索按钮
        if hasattr(ui, 'gene_search_btn'):
            ui.gene_search_btn.clicked.connect(self.func.search_gene)

        # 回车触发搜索
        if hasattr(ui, 'gene_search_input'):
            ui.gene_search_input.returnPressed.connect(self.func.search_gene)

    def set_volume(self, value):
        """设置音量（供 fix_music_controller_bindings 绑定音量滑块）"""
        from script.mods_layer.mod_manager import global_mod_manager
        mod_instance = global_mod_manager.get_current_mod()
        if hasattr(mod_instance, 'global_music_player'):
            mod_instance.global_music_player.set_volume(value / 100.0)

        if hasattr(self.parent, '_sync_all_volume_sliders_from_subinterface'):
            self.parent._sync_all_volume_sliders_from_subinterface(value)

    # ========== 信号处理 ==========

    def on_filter_mode_changed(self):
        """基因筛选模式切换：显示/隐藏对应控件"""
        ui = self.sc_monocle_genelists_ui
        mode = ui.gene_filter_combo.currentData()

        # 高变基因N输入框
        show_hvg = (mode == 'hvg')
        ui.hvg_n_label.setVisible(show_hvg)
        ui.hvg_n_spinbox.setVisible(show_hvg)

        # 外部基因列表导入
        show_list = (mode == 'list')
        ui.gene_list_label.setVisible(show_list)
        ui.btn_import_gene_list.setVisible(show_list)
        ui.gene_list_path_label.setVisible(show_list)

    def import_gene_list(self):
        """导入外部基因列表文件"""
        file_path = self.func.get_open_file_path(
            caption="选择基因列表文件",
            filter_str="基因列表文件 (*.txt *.xlsx);;文本文件 (*.txt);;Excel文件 (*.xlsx);;所有文件 (*.*)"
        )

        if not file_path:
            return

        # 验证文件可读取
        success, genes, err_msg = self.analysis.load_gene_list_from_file(file_path)
        if not success:
            self.func.alert_error(err_msg)
            self.func.log(f"❌ 导入基因列表失败: {err_msg}")
            return

        self._imported_gene_list_file = file_path
        file_name = os.path.basename(file_path)
        self.sc_monocle_genelists_ui.gene_list_path_label.setText(
            f"{file_name} ({len(genes)}个基因)"
        )
        self.func.log(f"✓ 导入基因列表: {file_path}")
        self.func.log(f"  共 {len(genes)} 个基因")

    def run_graph_test(self):
        """运行graph_test（QThread后台执行）"""
        # 同步共享数据
        self._sync_shared_data_from_main_window()

        if not self.analysis.has_shared_data():
            self.func.alert_error(
                "未共享CDS rds路径\n\n请先在「数据加载类」页面加载rds文件"
            )
            return

        # 获取筛选模式
        ui = self.sc_monocle_genelists_ui
        mode = ui.gene_filter_combo.currentData()

        # 参数校验
        if mode == 'list' and not self._imported_gene_list_file:
            self.func.alert_error("请先导入基因列表文件")
            return

        n_top = int(ui.hvg_n_spinbox.value()) if hasattr(ui.hvg_n_spinbox, 'value') else 2000

        # 更新数据信息
        self.func.update_data_info({
            'dataset': self.analysis.dataset_name or '未知',
            'rds_path': self.analysis.cds_rds_path,
            'filter_mode': ui.gene_filter_combo.currentText(),
        })

        # 禁用运行按钮，显示进度条
        self.func.set_run_button_enabled(False)
        self.func.set_progress_visible(True)
        self.func.set_progress_value(5)
        self.func.log("=" * 50)
        self.func.log(f"开始运行graph_test")
        self.func.log(f"  数据集: {self.analysis.dataset_name}")
        self.func.log(f"  筛选模式: {ui.gene_filter_combo.currentText()}")
        if mode == 'hvg':
            self.func.log(f"  高变基因数N: {n_top}")
        elif mode == 'list':
            self.func.log(f"  基因列表文件: {self._imported_gene_list_file}")

        # 创建Worker和Thread
        self._worker = GraphTestWorker(
            analysis=self.analysis,
            gene_filter_mode=mode,
            n_top=n_top,
            gene_list_file=self._imported_gene_list_file
        )
        self._thread = QThread()
        self._worker.moveToThread(self._thread)

        # 信号连接
        self._thread.started.connect(self._worker.run)
        self._worker.progress.connect(self._on_graph_test_progress)
        self._worker.finished.connect(self._on_graph_test_finished)
        self._worker.finished.connect(self._thread.quit)
        self._thread.finished.connect(self._worker.deleteLater)
        self._thread.finished.connect(self._thread.deleteLater)

        # 启动线程
        self._thread.start()

    def _on_graph_test_progress(self, msg):
        """graph_test进度回调（主线程）"""
        self.func.log(msg)
        # 模拟进度推进（R脚本无精确进度，按时间递增）
        current = self.sc_monocle_genelists_ui.progress_bar.value()
        if current < 90:
            self.func.set_progress_value(current + 2)

    def _on_graph_test_finished(self, success, result):
        """graph_test完成回调（主线程）

        阶段一：只填充所有基因表格，不自动筛选
        阶段二：由用户点击"重新筛显著基因"按钮触发（apply_threshold）
        """
        self.func.set_progress_visible(False)
        self.func.set_run_button_enabled(True)

        if not success:
            self.func.alert_failure(f"graph_test执行失败:\n{result}")
            self.func.log(f"❌ graph_test失败: {result}")
            return

        # 阶段一：只填充所有基因表格
        all_df = result
        self.func.fill_all_genes_table(all_df)

        # 清空显著基因表格（阶段二未执行）
        self.func.fill_significant_table(None)

        # 更新统计信息：总基因数显示，显著基因数为0（待用户筛选）
        self.func.update_stats_info(all_df, None)
        self.func.set_filter_button_enabled(True)
        # 阶段一完成后启用阶段三按钮（计算范围=全部基因时可用）
        self.func.set_stage3_button_enabled(True)

        self.func.alert_success(
            f"阶段一 graph_test 完成!\n\n"
            f"总基因数: {len(all_df)}\n\n"
            f"请调整阈值后点击「筛选显著基因」按钮执行阶段二筛选"
        )
        self.func.log(f"✓ 阶段一完成: 共 {len(all_df)} 个基因")
        self.func.log(f"  阶段二：调整阈值字段和阈值后点击「筛选显著基因」")

        # 切换到所有基因列表标签页
        self.sc_monocle_genelists_ui.genelists_tabs.setCurrentIndex(0)

    def apply_threshold(self):
        """阶段二：重新筛显著基因（不重跑R脚本）

        从缓存的 graph_test 结果 DataFrame 中按 p/q 阈值筛选
        纯 Python pandas 操作，不调用 R
        筛选完成后切换到显著基因列表标签页
        """
        threshold_field = self.sc_monocle_genelists_ui.threshold_field_combo.currentData()
        threshold_value = self.sc_monocle_genelists_ui.threshold_value_edit.text()

        if not threshold_field:
            self.func.alert_error("请选择阈值字段")
            return

        if not threshold_value:
            self.func.alert_error("请输入阈值")
            return

        # 检查是否有缓存结果
        cached_df = self.analysis.get_cached_result()
        if cached_df is None:
            self.func.alert_error("无 graph_test 结果，请先运行阶段一")
            self.func.log("❌ 筛选失败: 缓存结果为空，请先运行阶段一")
            return

        # 诊断日志：打印当前缓存DataFrame信息
        self.func.log("=" * 50)
        self.func.log("阶段二筛选开始（纯Python pandas操作，不调用R）")
        self.func.log(f"  阈值字段: {threshold_field}")
        self.func.log(f"  阈值: {threshold_value}")
        self.func.log(f"  缓存结果行数: {len(cached_df)}")
        self.func.log(f"  缓存结果列名: {list(cached_df.columns)}")

        try:
            success, result, is_error = self.analysis.filter_significant_genes(
                threshold_field=threshold_field,
                threshold_value=threshold_value
            )
        except Exception as e:
            import traceback
            err_detail = f"{e}\n{traceback.format_exc()}"
            self.func.alert_error(f"筛选异常: {e}")
            self.func.log(f"❌ 筛选异常: {err_detail}")
            return

        if not success:
            # 注意：失败时 result 是错误消息字符串
            self.func.alert_error(f"筛选失败: {result}")
            self.func.log(f"❌ 筛选失败: {result}")
            return

        sig_df = result
        all_df = cached_df
        self.func.fill_significant_table(sig_df)
        self.func.update_stats_info(all_df, sig_df)

        self.func.log(
            f"✓ 阶段二筛选完成: {threshold_field} < {threshold_value}, "
            f"显著 {len(sig_df)} / {len(all_df)} 基因"
        )
        self.func.alert_success(
            f"阶段二筛选完成!\n\n"
            f"阈值: {threshold_field} < {threshold_value}\n"
            f"显著基因数: {len(sig_df)} / {len(all_df)}"
        )

        # 切换到显著基因列表标签页
        self.sc_monocle_genelists_ui.genelists_tabs.setCurrentIndex(1)

    # 阶段三：基因表达与伪时间关系分析（上下调分类） ==========

    def run_spearman_analysis(self):
        """阶段三：运行基因表达与伪时间关系分析（QThread后台执行）

        支持三种算法：Spearman秩相关、LOESS局部加权回归、Wilcoxon检验
        完成后自动分类上下调（rho>0且p<阈值=up, rho<0且p<阈值=down）
        """
        # 同步共享数据
        self._sync_shared_data_from_main_window()

        if not self.analysis.has_shared_data():
            self.func.alert_error(
                "未共享CDS rds路径\n\n请先在「数据加载类」页面加载rds文件"
            )
            return

        # 检查阶段一缓存结果
        cached_df = self.analysis.get_cached_result()
        if cached_df is None:
            self.func.alert_error("无graph_test结果，请先运行阶段一")
            return

        # 获取参数
        ui = self.sc_monocle_genelists_ui
        algorithm = ui.stage3_algorithm_combo.currentData()
        calc_scope = ui.calc_scope_combo.currentData()
        threshold_field = ui.stage3_threshold_field_combo.currentData()
        threshold_value = ui.stage3_threshold_value_edit.text()

        if not threshold_value:
            self.func.alert_error("请输入分类阈值")
            return

        # 如果计算范围是显著基因，检查是否有显著基因
        if calc_scope == 'sig':
            sig_df = self.func._sig_genes_df
            if sig_df is None or len(sig_df) == 0:
                self.func.alert_error(
                    "计算范围为「显著基因」但阶段二未运行或无显著基因\n\n"
                    "请先运行阶段二筛选显著基因，或切换计算范围为「全部基因」"
                )
                return

        # 更新数据信息
        self.func.update_data_info({
            'dataset': self.analysis.dataset_name or '未知',
            'rds_path': self.analysis.cds_rds_path,
            'filter_mode': f"阶段三: {algorithm} + {ui.calc_scope_combo.currentText()}",
        })

        # 禁用运行按钮，显示进度条
        self.func.set_run_button_enabled_stage3(False)
        self.func.set_progress_visible(True)
        self.func.set_progress_value(5)
        self.func.log("=" * 50)
        self.func.log(f"开始运行阶段三 {algorithm}分析")
        self.func.log(f"  数据集: {self.analysis.dataset_name}")
        self.func.log(f"  算法: {ui.stage3_algorithm_combo.currentText()}")
        self.func.log(f"  计算范围: {ui.calc_scope_combo.currentText()}")
        self.func.log(f"  分类阈值字段: {threshold_field}")
        self.func.log(f"  分类阈值: {threshold_value}")

        # 创建Worker和Thread
        self._spearman_worker = SpearmanWorker(
            analysis=self.analysis,
            calc_scope=calc_scope,
            threshold_field=threshold_field,
            threshold_value=threshold_value,
            algorithm=algorithm
        )
        self._spearman_thread = QThread()
        self._spearman_worker.moveToThread(self._spearman_thread)

        # 信号连接
        self._spearman_thread.started.connect(self._spearman_worker.run)
        self._spearman_worker.progress.connect(self._on_spearman_progress)
        self._spearman_worker.finished.connect(self._on_spearman_finished)
        self._spearman_worker.finished.connect(self._spearman_thread.quit)
        self._spearman_thread.finished.connect(self._spearman_worker.deleteLater)
        self._spearman_thread.finished.connect(self._spearman_thread.deleteLater)

        # 启动线程
        self._spearman_thread.start()

    def _on_spearman_progress(self, msg):
        """Spearman分析进度回调（主线程）"""
        self.func.log(msg)
        # 模拟进度推进（R脚本无精确进度，按时间递增）
        current = self.sc_monocle_genelists_ui.progress_bar.value()
        if current < 90:
            self.func.set_progress_value(current + 2)

    def _on_spearman_finished(self, success, result):
        """Spearman分析完成回调（主线程）

        阶段三：运行完成后自动调用 classify_up_down_genes 分类上下调
        """
        self.func.set_progress_visible(False)
        self.func.set_run_button_enabled_stage3(True)

        if not success:
            self.func.alert_failure(f"阶段三分析失败:\n{result}")
            self.func.log(f"❌ 阶段三失败: {result}")
            return

        # 阶段三R脚本完成，自动分类上下调
        ui = self.sc_monocle_genelists_ui
        threshold_field = ui.stage3_threshold_field_combo.currentData()
        threshold_value = ui.stage3_threshold_value_edit.text()

        self.func.log("=" * 50)
        self.func.log("阶段三分类开始（纯Python pandas操作，不调用R）")
        self.func.log(f"  分类阈值字段: {threshold_field}")
        self.func.log(f"  分类阈值: {threshold_value}")

        try:
            success, result_dict, is_error = self.analysis.classify_up_down_genes(
                threshold_field=threshold_field,
                threshold_value=threshold_value
            )
        except Exception as e:
            import traceback
            err_detail = f"{e}\n{traceback.format_exc()}"
            self.func.alert_error(f"分类异常: {e}")
            self.func.log(f"❌ 分类异常: {err_detail}")
            return

        if not success:
            # 失败时 result_dict 是错误消息字符串
            self.func.alert_error(f"上下调分类失败: {result_dict}")
            self.func.log(f"❌ 上下调分类失败: {result_dict}")
            return

        # 填充3个表格
        self.func.fill_stage3_tables(result_dict)

        all_df = result_dict['all']
        up_df = result_dict['up']
        down_df = result_dict['down']

        # 更新统计信息（添加上下调基因数）
        self.func.update_stats_info(all_df, None, stage3_df=result_dict['all'])

        self.func.log(
            f"✓ 阶段三完成: 总体 {len(all_df)} 基因, "
            f"上调 {len(up_df)}, 下调 {len(down_df)}"
        )
        self.func.alert_success(
            f"阶段三 Spearman分析完成!\n\n"
            f"计算基因数: {len(all_df)}\n"
            f"上调基因: {len(up_df)}\n"
            f"下调基因: {len(down_df)}\n\n"
            f"阈值: {threshold_field} < {threshold_value}"
        )

        # 切换到上下调总体标签页（索引2）
        self.sc_monocle_genelists_ui.genelists_tabs.setCurrentIndex(2)

        # 阶段三完成后启用阶段四按钮和导出按钮
        self.func.set_stage4_button_enabled(True)
        self.func.set_export_button_enabled(True)

    # ========== 阶段四：火山图可视化 ==========

    def run_volcano_plot(self):
        """阶段四：绘制火山图（基于阶段三结果）"""
        # 同步共享数据
        self._sync_shared_data_from_main_window()

        if not self.analysis.has_shared_data():
            self.func.alert_error(
                "未共享CDS rds路径\n\n请先在「数据加载类」页面加载rds文件"
            )
            return

        # 检查阶段三缓存结果
        cached_df = self.analysis.get_cached_spearman_result()
        if cached_df is None:
            self.func.alert_error("无阶段三结果，请先运行阶段三")
            return

        # 获取参数
        ui = self.sc_monocle_genelists_ui
        x_axis = ui.stage4_x_axis_combo.currentData()
        y_axis = ui.stage4_y_axis_combo.currentData()
        fc_threshold = ui.stage4_fc_threshold_edit.text()
        p_threshold = ui.stage4_p_threshold_edit.text()
        top_n = int(ui.stage4_top_n_spinbox.value()) if hasattr(ui.stage4_top_n_spinbox, 'value') else 10

        # 参数校验
        try:
            fc_threshold = float(fc_threshold)
        except (ValueError, TypeError):
            self.func.alert_error("FC阈值必须是数字")
            return

        try:
            p_threshold = float(p_threshold)
        except (ValueError, TypeError):
            self.func.alert_error("p/q值阈值必须是数字")
            return

        # 更新数据信息
        self.func.update_data_info({
            'dataset': self.analysis.dataset_name or '未知',
            'rds_path': self.analysis.cds_rds_path,
            'filter_mode': f"阶段四: 火山图",
        })

        # 禁用运行按钮，显示进度条
        self.func.set_run_button_enabled_stage4(False)
        self.func.set_progress_visible(True)
        self.func.set_progress_value(5)
        self.func.log("=" * 50)
        self.func.log(f"开始运行阶段四 火山图可视化")
        self.func.log(f"  数据集: {self.analysis.dataset_name}")
        self.func.log(f"  X轴: {ui.stage4_x_axis_combo.currentText()}")
        self.func.log(f"  Y轴: {ui.stage4_y_axis_combo.currentText()}")
        self.func.log(f"  FC阈值: {fc_threshold}")
        self.func.log(f"  p/q阈值: {p_threshold}")
        self.func.log(f"  标记基因数: {top_n}")

        # 直接调用（火山图通常很快，不需要QThread）
        success, result, is_error = self.analysis.run_volcano_plot(
            x_axis=x_axis,
            y_axis=y_axis,
            fc_threshold=fc_threshold,
            p_threshold=p_threshold,
            top_n=top_n,
            progress_callback=self._on_volcano_progress
        )

        self.func.set_progress_visible(False)
        self.func.set_run_button_enabled_stage4(True)

        if not success:
            self.func.alert_failure(f"火山图绘制失败:\n{result}")
            self.func.log(f"❌ 阶段四失败: {result}")
            return

        # 显示火山图
        image_path = result
        self.func.show_volcano_image(image_path)
        self.func.log(f"✓ 阶段四完成: 火山图已保存并显示")
        self.func.alert_success(
            f"阶段四 火山图绘制完成!\n\n"
            f"图片路径: {image_path}"
        )

    def export_xlsx(self):
        """导出基因列表分析结果到xlsx文件（包含所有阶段数据）"""
        self.func.log("=" * 60)
        self.func.log("[DEBUG] 导出Excel开始")
        try:
            cached_stage3_df = self.analysis.get_cached_stage3_result()
            cached_graph_test_df = self.analysis.get_cached_result()
            
            all_genes_df = self.func._all_genes_df
            sig_genes_df = self.func._sig_genes_df
            
            self.func.log(f"[DEBUG] cached_stage3_df长度: {len(cached_stage3_df) if cached_stage3_df is not None else 0}")
            self.func.log(f"[DEBUG] cached_graph_test_df长度: {len(cached_graph_test_df) if cached_graph_test_df is not None else 0}")
            self.func.log(f"[DEBUG] all_genes_df长度: {len(all_genes_df) if all_genes_df is not None else 0}")
            self.func.log(f"[DEBUG] sig_genes_df长度: {len(sig_genes_df) if sig_genes_df is not None else 0}")
            
            if cached_stage3_df is None or len(cached_stage3_df) == 0:
                self.func.alert_error("请先运行阶段三分析")
                self.func.log("[DEBUG] 阶段三缓存为空，返回")
                return

            try:
                import openpyxl
                from openpyxl.utils.dataframe import dataframe_to_rows
                self.func.log("[DEBUG] openpyxl导入成功")
            except ImportError as e:
                self.func.alert_error("请安装 openpyxl 库以导出xlsx文件")
                self.func.log(f"[DEBUG] openpyxl导入失败: {e}")
                return

            ui = self.sc_monocle_genelists_ui
            
            gene_filter_mode = getattr(ui, 'gene_filter_combo', None)
            gene_filter_mode = gene_filter_mode.currentText() if gene_filter_mode else "未知"
            
            hvg_n = getattr(ui, 'hvg_n_spinbox', None)
            hvg_n = hvg_n.value() if hvg_n else 0
            
            stage2_threshold_field = getattr(ui, 'threshold_field_combo', None)
            stage2_threshold_field = stage2_threshold_field.currentText() if stage2_threshold_field else "未知"
            
            stage2_threshold_value = getattr(ui, 'threshold_value_edit', None)
            stage2_threshold_value = stage2_threshold_value.text() if stage2_threshold_value else "未知"
            
            stage3_algorithm = getattr(ui, 'stage3_algorithm_combo', None)
            stage3_algorithm = stage3_algorithm.currentText() if stage3_algorithm else "未知"
            
            stage3_threshold_field = getattr(ui, 'stage3_threshold_field_combo', None)
            stage3_threshold_field = stage3_threshold_field.currentText() if stage3_threshold_field else "未知"
            
            stage3_threshold_value = getattr(ui, 'stage3_threshold_value_edit', None)
            stage3_threshold_value = stage3_threshold_value.text() if stage3_threshold_value else "未知"
            
            dataset_name = self.analysis.dataset_name or "未知数据集"

            default_name = f"基因列表分析_{dataset_name}.xlsx"
            save_path = self.func.get_save_file_path("导出基因列表分析结果", default_name, "Excel文件 (*.xlsx)")

            if save_path:
                wb = openpyxl.Workbook()
                
                ws1 = wb.active
                ws1.title = "统计信息"
                
                if 'direction' in cached_stage3_df.columns:
                    up_count = len(cached_stage3_df[cached_stage3_df['direction'] == 'up'])
                    down_count = len(cached_stage3_df[cached_stage3_df['direction'] == 'down'])
                    stable_count = len(cached_stage3_df[cached_stage3_df['direction'] == 'not_significant'])
                else:
                    up_count = 0
                    down_count = 0
                    stable_count = 0
                
                total_stage3_count = len(cached_stage3_df)
                total_graph_test_count = len(cached_graph_test_df) if cached_graph_test_df is not None else 0
                sig_count = len(sig_genes_df) if sig_genes_df is not None else 0
                
                stats_data = [
                    ["统计项", "数值"],
                    ["数据集名称", dataset_name],
                    ["阶段一 基因筛选模式", gene_filter_mode],
                    ["阶段一 高变基因数N", hvg_n],
                    ["阶段一 总体基因数", total_graph_test_count],
                    ["阶段二 筛选阈值字段", stage2_threshold_field],
                    ["阶段二 筛选阈值", stage2_threshold_value],
                    ["阶段二 显著基因数", sig_count],
                    ["阶段三 分析算法", stage3_algorithm],
                    ["阶段三 分类阈值字段", stage3_threshold_field],
                    ["阶段三 分类阈值", stage3_threshold_value],
                    ["阶段三 总体基因数", total_stage3_count],
                    ["阶段三 上调基因数", up_count],
                    ["阶段三 下调基因数", down_count],
                    ["阶段三 稳定基因数", stable_count],
                ]
                for row in stats_data:
                    ws1.append(row)
                
                if cached_graph_test_df is not None:
                    ws2 = wb.create_sheet(title="阶段一 graph_test结果")
                    for r in dataframe_to_rows(cached_graph_test_df, index=False, header=True):
                        ws2.append(r)
                
                if sig_genes_df is not None:
                    ws3 = wb.create_sheet(title="阶段二 显著基因列表")
                    for r in dataframe_to_rows(sig_genes_df, index=False, header=True):
                        ws3.append(r)
                
                ws4 = wb.create_sheet(title="阶段三 总体基因列表")
                for r in dataframe_to_rows(cached_stage3_df, index=False, header=True):
                    ws4.append(r)
                
                ws5 = wb.create_sheet(title="阶段三 上调基因列表")
                if 'direction' in cached_stage3_df.columns:
                    up_df = cached_stage3_df[cached_stage3_df['direction'] == 'up']
                else:
                    up_df = cached_stage3_df.head(0)
                for r in dataframe_to_rows(up_df, index=False, header=True):
                    ws5.append(r)
                
                ws6 = wb.create_sheet(title="阶段三 下调基因列表")
                if 'direction' in cached_stage3_df.columns:
                    down_df = cached_stage3_df[cached_stage3_df['direction'] == 'down']
                else:
                    down_df = cached_stage3_df.head(0)
                for r in dataframe_to_rows(down_df, index=False, header=True):
                    ws6.append(r)
                
                wb.save(save_path)
                self.func.alert_success(f"结果已保存到:\n{save_path}")
                self.func.log(f"✓ 导出完成: {save_path}")
        except Exception as e:
            import traceback
            err_detail = f"{e}\n{traceback.format_exc()}"
            self.func.alert_failure(f"导出失败: {str(e)}")
            self.func.log(f"❌ 导出失败: {err_detail}")

    def _on_volcano_progress(self, msg):
        """火山图进度回调"""
        self.func.log(msg)
        current = self.sc_monocle_genelists_ui.progress_bar.value()
        if current < 90:
            self.func.set_progress_value(current + 5)

    def _on_x_axis_changed(self, index):
        """X轴指标切换时自动调整FC阈值"""
        ui = self.sc_monocle_genelists_ui
        if not hasattr(ui, 'stage4_x_axis_combo') or not hasattr(ui, 'stage4_fc_threshold_edit'):
            return
        
        x_axis = ui.stage4_x_axis_combo.currentData()
        if x_axis == 'rho':
            ui.stage4_fc_threshold_edit.setText("0.2")
        elif x_axis == 'log2fc':
            ui.stage4_fc_threshold_edit.setText("1.0")

    # ========== 数据同步 ==========

    def _sync_shared_data_from_main_window(self):
        """从main_window读取共享的cds_rds_path和dataset_name"""
        try:
            shared_path = getattr(self.parent, 'shared_monocle_cds_rds_path', None)
            shared_dataset = getattr(self.parent, 'shared_monocle_dataset_name', None)

            if shared_path and shared_dataset:
                self.analysis.set_cds_rds_path(shared_path, shared_dataset)
                self.func.log(f"✓ 已同步CDS路径: {shared_dataset}")
        except Exception as e:
            print(f"同步共享数据失败: {e}")

    def sync_data_from_single_cell_main(self, single_cell_bind=None):
        """从scRNAseq主页同步数据（保留接口，基因列表类依赖数据加载类）"""
        try:
            # 优先同步共享数据
            self._sync_shared_data_from_main_window()
        except Exception as e:
            print(f"Monocle基因列表类同步数据时出错: {str(e)}")

    def on_page_activated(self):
        """页面激活时调用（由主容器在切换到本页面时调用）"""
        self._sync_shared_data_from_main_window()

        # 更新数据信息显示
        if self.analysis.has_shared_data():
            self.func.update_data_info({
                'dataset': self.analysis.dataset_name,
                'rds_path': self.analysis.cds_rds_path,
                'filter_mode': self.sc_monocle_genelists_ui.gene_filter_combo.currentText(),
            })
        else:
            self.func.update_data_info({
                'dataset': '未共享',
                'rds_path': '请先在数据加载类页面加载rds',
                'filter_mode': self.sc_monocle_genelists_ui.gene_filter_combo.currentText(),
            })
